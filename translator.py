"""
Local live speech translator (Japanese <-> English) with an HTML/CSS user interface.

Pipeline:  microphone -> phrase segmentation -> language ID (Whisper-base)
           -> speech recognition (LFM2.5-Audio) -> translation (LFM2.5-1.2B-JP) -> UI

Threads:
  - sounddevice callback : pushes audio blocks into audio_q
  - segmenter            : cuts the audio stream into phrases at pauses -> seg_q
  - worker               : phrase -> source text -> translation -> ui_q
  - broadcaster + HTTP   : streams events to the window (Server-Sent Events)
  - main thread          : shows the window (pywebview, or Edge in app mode)
"""
import collections
import json
import os
import sys

# Started without a console (pythonw / the .vbs launcher)? Then stdout/stderr don't exist
# and libraries that print progress bars would crash - send everything to app.log instead.
if sys.stdout is None or sys.stderr is None:
    _log = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "app.log"),
                "w", encoding="utf-8", buffering=1)
    sys.stdout = sys.stdout or _log
    sys.stderr = sys.stderr or _log

import queue
import re
import subprocess
import threading
import time
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import numpy as np
import sounddevice as sd
import torch
from liquid_audio import LFM2AudioModel, LFM2AudioProcessor, ChatState
from transformers import (AutoModelForCausalLM, AutoTokenizer, TextStreamer,
                          WhisperForConditionalGeneration, WhisperProcessor)

# ----------------------------- SETTINGS -----------------------------
ASR_JA_REPO = "LiquidAI/LFM2.5-Audio-1.5B-JP"  # speech -> Japanese text
ASR_EN_REPO = "LiquidAI/LFM2.5-Audio-1.5B"     # speech -> English text
PROMPT_JA = "Perform ASR in japanese."         # fixed system prompt of the JP model
PROMPT_EN = "Perform ASR."                     # system prompt of the base (English) model
LID_REPO = "openai/whisper-base"               # only used to detect the spoken language
MT_REPO = "LiquidAI/LFM2.5-1.2B-JP"            # translation JP <-> EN
# Faster dedicated 350M translation model - uncomment if 1.2B is too slow:
# MT_REPO = "LiquidAI/LFM2-350M-ENJP-MT"
MT_STYLE = "enjp-mt" if "ENJP-MT" in MT_REPO else "chat"

MIC_DEVICE = None        # microphone index from `python -m sounddevice`; None = system default
SR = 16000               # sample rate
BLOCK = int(0.1 * SR)    # one audio block = 0.1 s
THRESHOLD = 0.005        # initial volume threshold (adjustable in the UI)
SILENCE_BLOCKS = 7       # this many quiet blocks (~0.7 s) end a phrase
MAX_BLOCKS = 150         # hard limit of ~15 s per phrase
MIN_BLOCKS = 5           # ignore phrases shorter than ~0.5 s (clicks, noise)
PARTIAL_EVERY_BLOCKS = 20  # live (interim) recognition every ~2 s while speaking
PREROLL_BLOCKS = 3       # keep ~0.3 s before speech starts so the first syllable isn't cut
DEVICE = "cuda" if torch.cuda.is_available() else "cpu"
# ---------------------------------------------------------------------

audio_q = queue.Queue()             # raw audio blocks
seg_q = queue.Queue()               # finished phrases
ui_q = queue.Queue()                # events for the window
text_q = queue.Queue()              # translation requests from the "Text" tab
partial_q = queue.Queue(maxsize=1)  # newest unfinished chunk for interim recognition
listening = threading.Event()
shutdown = threading.Event()
state = {"mode": "auto",        # spoken language: "auto" | "ja" | "en"
         "threshold": THRESHOLD,
         "level": 0.0,          # current microphone level
         "partials": True,      # show text while speaking
         "stream_tr": False,    # show the translation while it is being generated
         "transcript": "txt",   # save a transcript on Stop: "off" | "txt" | "docx"
         "busy": False}         # the worker is processing a phrase right now


def audio_callback(indata, frames, time_info, status):
    block = indata[:, 0]
    state["level"] = float(np.sqrt(np.mean(block ** 2)))
    if listening.is_set():
        audio_q.put(block.copy())


def normalize_audio(wav: np.ndarray, target_peak: float = 0.9, max_gain: float = 30.0):
    """Boosts quiet recordings (up to 30x) so the model hears them better."""
    wav = wav - float(np.mean(wav))
    peak = float(np.max(np.abs(wav)))
    if peak < 1e-4:
        return wav.astype(np.float32)
    return (wav * min(target_peak / peak, max_gain)).astype(np.float32)


def put_latest(q, item):
    """Puts an item into a maxsize=1 queue, replacing the old one."""
    try:
        q.get_nowait()
    except queue.Empty:
        pass
    try:
        q.put_nowait(item)
    except queue.Full:
        pass


SENT_RE = re.compile(r"(?<=[。！？!?])\s*|(?<=[.])\s+")


def split_sentences(text: str):
    return [p.strip() for p in SENT_RE.split(text) if p and p.strip()]


# ------------------------- transcript of a session -------------------------
session = {"start": None, "entries": []}   # a session = from Start to Stop
session_lock = threading.Lock()


def session_start():
    with session_lock:
        if session["start"] is None:
            session["start"], session["entries"] = time.time(), []


def session_add(src: str, orig: str, tr: str):
    with session_lock:
        if session["start"] is not None:
            session["entries"].append((time.time(), src, orig, tr))


def _clock(ts: float) -> str:
    return time.strftime("%H:%M:%S", time.localtime(ts))


def _header(start: float, n: int) -> str:
    secs = int(time.time() - start)
    return (f"Started: {time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(start))}   |   "
            f"Duration: {secs // 60} min {secs % 60} s   |   Phrases: {n}")


def write_txt(path: str, start: float, entries: list):
    with open(path, "w", encoding="utf-8") as f:
        f.write("Live Translator transcript\n" + _header(start, len(entries)) + "\n\n")
        for ts, src, orig, tr in entries:
            f.write(f"[{_clock(ts)}] {'JP' if src == 'Japanese' else 'EN'}\n{orig}\n-> {tr}\n\n")


def write_docx(path: str, start: float, entries: list):
    from docx import Document
    doc = Document()
    doc.add_heading("Live Translator transcript", level=1)
    doc.add_paragraph(_header(start, len(entries)))
    for ts, src, orig, tr in entries:
        p = doc.add_paragraph()
        p.add_run(f"[{_clock(ts)}] {'JP' if src == 'Japanese' else 'EN'}  ").bold = True
        p.add_run(orig).italic = True
        doc.add_paragraph(tr)
    doc.save(path)


def save_session(restart: bool = False):
    """Writes the current session to TRANSCRIPT_DIR. Returns the file path or None."""
    with session_lock:
        start, entries = session["start"], session["entries"]
        session["start"] = time.time() if restart else None
        session["entries"] = []
    fmt = state["transcript"]
    if start is None or not entries or fmt == "off":
        return None
    os.makedirs(TRANSCRIPT_DIR, exist_ok=True)
    stamp = time.strftime("%Y-%m-%d_%H-%M-%S", time.localtime(start))
    if fmt == "docx":
        try:
            path = os.path.join(TRANSCRIPT_DIR, f"transcript_{stamp}.docx")
            write_docx(path, start, entries)
            return path
        except ImportError:
            ui_q.put(("notice", "python-docx is not installed, so the transcript was saved "
                                "as TXT (pip install python-docx)"))
    path = os.path.join(TRANSCRIPT_DIR, f"transcript_{stamp}.txt")
    write_txt(path, start, entries)
    return path


def finish_session_async():
    """On Stop: wait (max 20 s) until pending phrases are translated, then save the transcript."""
    def run():
        t0 = time.time()
        while (not seg_q.empty() or state["busy"]) and time.time() - t0 < 20:
            time.sleep(0.2)
        try:
            path = save_session(restart=listening.is_set())
        except Exception as e:  # noqa: BLE001
            ui_q.put(("error", f"Could not save the transcript: {e}"))
            return
        if path:
            ui_q.put(("saved", path))
    threading.Thread(target=run, daemon=True).start()


PREFACE_RE = re.compile(r"^\s*(?:(?:sure|certainly|of course|okay|ok)[!.,]?\s*)?"
                        r"(?:here(?:'s| is)[^:\n]{0,80}:\s*)?", re.I)
NOTE_RE = re.compile(r"\n\s*\n|\n\s*\(?(?:note|translation note)\b", re.I)
BAD_RE = re.compile(r"the (?:japanese|english) (?:sentence|text|phrase)|incomplete or|possibly a typo|"
                    r"if you meant|i(?:'m| am) sorry|i (?:cannot|can't) |translation would be|"
                    r"here is the translation", re.I)


def clean_translation(s: str) -> str:
    """Removes 'Sure! Here is the translation:', trailing notes and wrapping quotes."""
    s = PREFACE_RE.sub("", s.strip())
    s = NOTE_RE.split(s)[0].strip()
    if len(s) >= 2 and s[0] in "\"\u201c" and s[-1] in "\"\u201d":
        s = s[1:-1].strip()
    return s


def segmenter():
    """Cuts the continuous audio stream into phrases at pauses."""
    preroll = collections.deque(maxlen=PREROLL_BLOCKS)
    buf, silent, speaking = [], 0, False
    while not shutdown.is_set():
        try:
            block = audio_q.get(timeout=0.2)
        except queue.Empty:
            if not listening.is_set():  # listening was stopped mid-phrase: drop the leftovers
                buf, silent, speaking = [], 0, False
            continue

        rms = float(np.sqrt(np.mean(block ** 2)))
        loud = rms > state["threshold"]
        if not speaking:
            if loud:
                speaking, silent = True, 0
                buf = list(preroll) + [block]
            else:
                preroll.append(block)
            continue

        buf.append(block)
        
        silent = 0 if rms > state["threshold"] * 0.6 else silent + 1
        if (state["partials"] and silent < SILENCE_BLOCKS
                and len(buf) % PARTIAL_EVERY_BLOCKS == 0):
            put_latest(partial_q, np.concatenate(buf))
        if silent >= SILENCE_BLOCKS or len(buf) >= MAX_BLOCKS:
            try:
                partial_q.get_nowait()
            except queue.Empty:
                pass
            if len(buf) - silent >= MIN_BLOCKS:
                seg_q.put(np.concatenate(buf))
                ui_q.put(("queue", seg_q.qsize()))
            buf, silent, speaking = [], 0, False
            preroll.clear()


SPECIAL_RE = re.compile(r"<\|[^|>]*\|>")


def clean_text(s: str) -> str:
    """Removes special tokens (<|im_end|> etc.) and truncated characters."""
    return SPECIAL_RE.sub("", s).rstrip("\ufffd").strip()


def gpu_report(tag: str):
    if DEVICE == "cuda":
        free, total = torch.cuda.mem_get_info()
        print(f"[gpu] {tag}: PyTorch holds {torch.cuda.memory_allocated() / 2**30:.1f} GB | "
              f"free {free / 2**30:.1f} of {total / 2**30:.1f} GB", flush=True)


def load_lid():
    """Tiny Whisper model that only detects which language is being spoken."""
    proc = WhisperProcessor.from_pretrained(LID_REPO)
    model = WhisperForConditionalGeneration.from_pretrained(LID_REPO).eval().to(DEVICE)
    tk = proc.tokenizer
    return {"proc": proc, "model": model, "dev": DEVICE,
            "sot": tk.convert_tokens_to_ids("<|startoftranscript|>"),
            "ja": tk.convert_tokens_to_ids("<|ja|>"),
            "en": tk.convert_tokens_to_ids("<|en|>")}


def detect_speech_lang(lid, wav: np.ndarray) -> str:
    """Compares how much the audio resembles Japanese vs English."""
    feats = lid["proc"](wav, sampling_rate=SR, return_tensors="pt").input_features
    feats = feats.to(lid["dev"])
    dec = torch.tensor([[lid["sot"]]], device=lid["dev"])
    with torch.inference_mode():
        logits = lid["model"](input_features=feats, decoder_input_ids=dec).logits[0, -1]
    return "Japanese" if logits[lid["ja"]] > logits[lid["en"]] else "English"


class CallbackStreamer(TextStreamer):
    """Calls on_partial(text_so_far) whenever the model has produced new words."""

    def __init__(self, tokenizer, on_partial):
        super().__init__(tokenizer, skip_prompt=True, skip_special_tokens=True)
        self.on_partial, self.acc = on_partial, ""

    def on_finalized_text(self, text: str, stream_end: bool = False):
        self.acc += text
        self.on_partial(self.acc.strip())


def load_asr(repo: str, device: str):
    """Returns (processor, model) of an audio model placed on the given device."""
    processor = LFM2AudioProcessor.from_pretrained(repo, device=device).eval()
    model = LFM2AudioModel.from_pretrained(repo, device=device).eval()
    return processor, model


def move_asr(asr, device: str):
    processor, model = asr
    processor.to(device)
    model.to(device)


def transcribe(asr, wav: np.ndarray, prompt: str, on_text=None, should_stop=None) -> str:
    """Recognizes one phrase. The prompt depends on the model (PROMPT_JA / PROMPT_EN).
    on_text(s) is called after every new token; should_stop() can abort the recognition."""
    processor, model = asr
    chat = ChatState(processor)
    chat.new_turn("system")
    chat.add_text(prompt)
    chat.end_turn()
    chat.new_turn("user")
    chat.add_audio(torch.from_numpy(wav).unsqueeze(0), SR)
    chat.end_turn()
    chat.new_turn("assistant")
    ids = []
    with torch.inference_mode():
        for t in model.generate_sequential(**chat, max_new_tokens=512):
            if should_stop and should_stop():
                break
            if t.numel() == 1:  # text tokens
                ids.append(int(t.reshape(-1)[0]))
                if on_text:
                    on_text(clean_text(processor.text.decode(ids)))
    # decode all tokens together so multi-byte (Japanese) characters aren't broken
    return clean_text(processor.text.decode(ids)) if ids else ""


def worker():
    """Loads the models, then processes phrases from the queues."""
    try:
        # GPU memory (8 GB) is too small for both audio models plus the translator,
        # and Windows then silently spills into slow shared memory. So only ONE audio
        # model lives on the GPU at a time; the other waits in RAM and is swapped in on demand.
        ui_q.put(("status", "Loading speech model (JP)…"))
        jp_asr = load_asr(ASR_JA_REPO, DEVICE)

        ui_q.put(("status", "Loading speech model (EN)…"))
        en_asr = load_asr(ASR_EN_REPO, "cpu")

        ui_q.put(("status", "Loading language detector…"))
        lid = load_lid()

        ui_q.put(("status", "Loading translation model…"))
        tok = AutoTokenizer.from_pretrained(MT_REPO)
        mt_model = AutoModelForCausalLM.from_pretrained(MT_REPO, dtype=torch.bfloat16)
        mt_model = mt_model.to(DEVICE).eval()
        gpu_report("after loading all models")
    except Exception as e:  # noqa: BLE001
        ui_q.put(("error", f"Failed to load models: {e}"))
        return

    asr_models = {"Japanese": (jp_asr, PROMPT_JA), "English": (en_asr, PROMPT_EN)}
    active = {"lang": "Japanese"}

    def lang_asr(src):
        """Returns (asr, prompt) for a language, swapping models on the GPU if needed."""
        if DEVICE == "cuda" and active["lang"] != src:
            t0 = time.time()
            move_asr(asr_models[active["lang"]][0], "cpu")
            torch.cuda.empty_cache()
            move_asr(asr_models[src][0], "cuda")
            active["lang"] = src
            print(f"[swap] speech model on GPU -> {src} ({time.time() - t0:.1f} s)", flush=True)
        return asr_models[src]

    def pick_lang(wav):
        mode = state["mode"]
        if mode == "ja":
            return "Japanese"
        if mode == "en":
            return "English"
        return detect_speech_lang(lid, wav)  # auto: decided from the sound, not the text

    def run_mt(text: str, src: str, dst: str, on_partial, cap: int, strict: bool) -> str:
        t0 = time.time()
        if MT_STYLE == "enjp-mt":
            msgs = [{"role": "system", "content": f"Translate to {dst}."},
                    {"role": "user", "content": text}]
        else:
            system = (f"You are a translation engine. Translate the user's text from {src} to {dst}. "
                      "Output ONLY the translation: no preface, no notes, no explanations, no quotation "
                      "marks. If the text is incomplete or contains errors, translate it as it is "
                      "without commenting on it.")
            if strict:
                system += " The text is a spoken fragment: translate that fragment literally, whatever it is."
            ex_src, ex_dst = ("こんにちは。", "Hello.") if src == "Japanese" else ("Hello.", "こんにちは。")
            msgs = [{"role": "system", "content": system},
                    {"role": "user", "content": ex_src}, {"role": "assistant", "content": ex_dst},
                    {"role": "user", "content": text}]
        inputs = tok.apply_chat_template(
            msgs, add_generation_prompt=True, return_tensors="pt",
            tokenize=True, return_dict=True,
        ).to(mt_model.device)
        # cap the answer length relative to the input so the model can't ramble on
        n_in = len(tok(text, add_special_tokens=False).input_ids)
        kwargs = dict(do_sample=False, max_new_tokens=min(cap, max(48, n_in * 4)))
        if on_partial:
            kwargs["streamer"] = CallbackStreamer(tok, on_partial)
        with torch.inference_mode():
            out = mt_model.generate(**inputs, **kwargs)
        n = inputs["input_ids"].shape[1]
        dt = time.time() - t0
        print(f"[translate] {out.shape[1] - n} tokens in {dt:.1f} s "
              f"({(out.shape[1] - n) / max(dt, 1e-6):.0f} tok/s) on {mt_model.device}", flush=True)
        return clean_translation(tok.decode(out[0][n:], skip_special_tokens=True))

    def translate(text: str, src: str, dst: str, on_partial=None, cap: int = 400) -> str:
        out = run_mt(text, src, dst, on_partial, cap, strict=False)
        if BAD_RE.search(out):  # the model commented instead of translating: retry once, stricter
            out = run_mt(text, src, dst, None, cap, strict=True)
            if BAD_RE.search(out):
                out = "[unclear phrase]"
        return out

    # warm-up: the very first GPU call is always slower, so do it before the user does
    try:
        ui_q.put(("status", "Warming up translation model…"))
        translate("こんにちは。", "Japanese", "English")
    except Exception as e:  # noqa: BLE001
        print("Warm-up failed:", e)

    ui_q.put(("ready", None))

    def live(s):
        ui_q.put(("live", s))

    while not shutdown.is_set():
        # 1) requests from the "Text" tab
        try:
            src, dst, text = text_q.get_nowait()
        except queue.Empty:
            pass
        else:
            try:
                ui_q.put(("text_result", translate(
                    text, src, dst, cap=1024,
                    on_partial=(lambda t: ui_q.put(("text_partial", t)))
                    if state["stream_tr"] else None)))
            except Exception as e:  # noqa: BLE001
                ui_q.put(("text_error", str(e)))
            continue

        # 2) finished phrases have priority over interim ones
        try:
            wav = seg_q.get_nowait()
        except queue.Empty:
            wav = None

        if wav is None:
            # 3) interim recognition of what is being said right now
            try:
                pwav = partial_q.get_nowait()
            except queue.Empty:
                time.sleep(0.05)
                continue
            if not (state["partials"] and listening.is_set()):
                continue
            try:
                pwav = normalize_audio(pwav)
                asr, prompt = lang_asr(pick_lang(pwav))
                transcribe(asr, pwav, prompt, on_text=live,
                           should_stop=lambda: (not seg_q.empty()) or shutdown.is_set())
            except Exception as e:  # noqa: BLE001
                ui_q.put(("error", f"Live recognition error: {e}"))
            continue

        try:
            state["busy"] = True
            ui_q.put(("busy", seg_q.qsize()))
            wav = normalize_audio(wav)
            t0 = time.time()
            src = pick_lang(wav)
            t1 = time.time()
            asr, prompt = lang_asr(src)
            t2 = time.time()
            text = transcribe(asr, wav, prompt, on_text=live)
            print(f"[asr] audio {len(wav) / SR:.1f} s | language ID {t1 - t0:.1f} s | "
                  f"model swap {t2 - t1:.1f} s | recognition {time.time() - t2:.1f} s | {src}",
                  flush=True)
            ui_q.put(("live_clear", None))
            if text:
                dst = "English" if src == "Japanese" else "Japanese"
                for sent in split_sentences(text):  # translate sentence by sentence
                    tr = translate(sent, src, dst,
                                   on_partial=(lambda t, _s=src, _o=sent:
                                               ui_q.put(("tr_live", (_s, _o, t))))
                                   if state["stream_tr"] else None)
                    ui_q.put(("pair", (src, sent, tr)))
                    session_add(src, sent, tr)
        except Exception as e:  # noqa: BLE001
            ui_q.put(("live_clear", None))
            ui_q.put(("error", f"Processing error: {e}"))
        finally:
            state["busy"] = False
            ui_q.put(("done", seg_q.qsize()))


JA_RE = re.compile(r"[\u3040-\u30ff\u31f0-\u31ff\u3400-\u4dbf\u4e00-\u9fff\uff66-\uff9f]")
LATIN_RE = re.compile(r"[A-Za-z]")


def detect_language(text: str):
    """Detects the language of typed text by script: Japanese (kana/kanji) vs Latin.
    Returns "Japanese", "English", or None if it can't be told."""
    ja = len(JA_RE.findall(text))
    lat = len(LATIN_RE.findall(text))
    if ja == 0 and lat == 0:
        return None
    # Japanese sentences often contain Latin words, so use the share instead of "any"
    return "Japanese" if ja / (ja + lat) >= 0.3 else "English"


# ============================ WEB UI SERVER ============================
HERE = os.path.dirname(os.path.abspath(__file__))
PORT = 8765
TRANSCRIPT_DIR = os.path.join(HERE, "transcripts")
ICON_PATH = os.path.join(HERE, "icon.ico")
subscribers = []                    # queues of connected windows (SSE)
sub_lock = threading.Lock()
clients = {"n": 0}
app_state = {"ready": False, "status": "Starting…"}


def publish(msg):
    with sub_lock:
        for q in subscribers:
            q.put(msg)


def broadcaster():
    """Forwards events from ui_q to the window, plus the mic level 10x per second."""
    last = 0.0
    while not shutdown.is_set():
        try:
            kind, data = ui_q.get(timeout=0.1)
        except queue.Empty:
            pass
        else:
            if kind == "ready":
                app_state["ready"] = True
            elif kind == "status":
                app_state["status"] = data
            publish({"type": kind, "data": data})
        now = time.time()
        if now - last > 0.1:
            last = now
            publish({"type": "level", "data": round(state["level"], 4)})


def handle_command(cmd: str, b: dict) -> dict:
    if cmd == "listen":
        if b.get("on"):
            session_start()
            listening.set()
        else:
            listening.clear()
            finish_session_async()
        return {"ok": True}
    if cmd == "open_folder":
        os.makedirs(TRANSCRIPT_DIR, exist_ok=True)
        if os.name == "nt":
            os.startfile(TRANSCRIPT_DIR)  # noqa: S606
        else:
            subprocess.Popen(["xdg-open", TRANSCRIPT_DIR])
        return {"ok": True}
    if cmd == "settings":
        if b.get("mode") in ("auto", "ja", "en"):
            state["mode"] = b["mode"]
        if "threshold" in b:
            state["threshold"] = float(b["threshold"])
        for key in ("partials", "stream_tr"):
            if key in b:
                state[key] = bool(b[key])
        if b.get("transcript") in ("off", "txt", "docx"):
            state["transcript"] = b["transcript"]
        return {"ok": True}
    if cmd == "translate":
        text = (b.get("text") or "").strip()
        if not text:
            return {"error": "Empty text"}
        if not app_state["ready"]:
            return {"error": "Models are still loading"}
        d = b.get("direction", "auto")
        if d == "ja-en":
            pair = ("Japanese", "English")
        elif d == "en-ja":
            pair = ("English", "Japanese")
        else:
            src = detect_language(text)
            if src is None:
                return {"error": "Couldn't detect the language (Japanese or English expected)"}
            pair = (src, "English" if src == "Japanese" else "Japanese")
        text_q.put((pair[0], pair[1], text))
        return {"ok": True, "direction": f"{pair[0]} → {pair[1]}"}
    return {"error": "unknown command"}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):  # keep the console quiet
        pass

    def _send(self, body: bytes, ctype: str, code: int = 200):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _sse(self, msg: dict):
        self.wfile.write(("data: " + json.dumps(msg, ensure_ascii=False) + "\n\n").encode("utf-8"))
        self.wfile.flush()

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            try:
                with open(os.path.join(HERE, "index.html"), "rb") as f:
                    self._send(f.read(), "text/html; charset=utf-8")
            except OSError:
                self._send(b"index.html not found next to translator.py",
                           "text/plain; charset=utf-8", 404)
        elif self.path == "/icon.ico":
            try:
                with open(ICON_PATH, "rb") as f:
                    self._send(f.read(), "image/x-icon")
            except OSError:
                self._send(b"", "image/x-icon", 404)
        elif self.path == "/events":
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream")
            self.send_header("Cache-Control", "no-cache")
            self.end_headers()
            q = queue.Queue()
            with sub_lock:
                subscribers.append(q)
                clients["n"] += 1
            try:
                self._sse({"type": "init", "data": {
                    "ready": app_state["ready"], "status": app_state["status"],
                    "listening": listening.is_set(), "mode": state["mode"],
                    "threshold": state["threshold"], "partials": state["partials"],
                    "transcript": state["transcript"], "stream_tr": state["stream_tr"]}})
                while not shutdown.is_set():
                    try:
                        msg = q.get(timeout=10)
                    except queue.Empty:
                        self.wfile.write(b": keepalive\n\n")
                        self.wfile.flush()
                        continue
                    self._sse(msg)
            except OSError:  # the window was closed
                pass
            finally:
                with sub_lock:
                    subscribers.remove(q)
                    clients["n"] -= 1
        else:
            self._send(b"not found", "text/plain", 404)

    def do_POST(self):
        n = int(self.headers.get("Content-Length") or 0)
        try:
            body = json.loads(self.rfile.read(n) or b"{}")
        except ValueError:
            body = {}
        result = handle_command(self.path.rsplit("/", 1)[-1], body)
        self._send(json.dumps(result, ensure_ascii=False).encode("utf-8"),
                   "application/json; charset=utf-8")


def start_server():
    for port in range(PORT, PORT + 20):
        try:
            srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
        except OSError:
            continue
        srv.daemon_threads = True
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        return srv, port
    raise RuntimeError("Could not start the local server (all ports are busy)")


def hard_exit(*_):
    """Terminates the process immediately (PyTorch/CUDA threads could keep it alive)."""
    try:  # window closed while listening: don't lose the transcript
        path = save_session()
        if path:
            print(f"Transcript saved: {path}", flush=True)
    except Exception:  # noqa: BLE001
        pass
    print("Application closed.", flush=True)
    os._exit(0)


def watchdog():
    """Exits once the window has been open and then had no connection for ~4 s."""
    seen, zero = False, 0
    while True:
        time.sleep(0.5)
        if clients["n"] > 0:
            seen, zero = True, 0
        elif seen:
            zero += 1
            if zero >= 8:
                hard_exit()


def wait_for_window_close():
    """Edge mode: return once the window is gone (its SSE connection dropped)."""
    seen, gone = False, None
    while not shutdown.is_set():
        time.sleep(0.5)
        if clients["n"] > 0:
            seen, gone = True, None
        elif seen:
            gone = gone or time.time()
            if time.time() - gone > 3:
                break


def apply_window_icon():
    """Best effort (Windows): show icon.ico in the title bar and the taskbar."""
    if os.name != "nt" or not os.path.exists(ICON_PATH):
        return
    import ctypes
    from ctypes import c_size_t, c_uint, c_void_p
    try:
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("LiveTranslator.App")
    except Exception:  # noqa: BLE001
        pass

    def run():
        user32 = ctypes.windll.user32
        user32.FindWindowW.restype = c_void_p
        user32.LoadImageW.restype = c_void_p
        user32.SendMessageW.argtypes = [c_void_p, c_uint, c_size_t, c_void_p]
        for _ in range(40):  # wait up to ~20 s for the window to appear
            hwnd = user32.FindWindowW(None, "Live Translator")
            if hwnd:
                for kind, size in ((1, 32), (0, 16)):  # ICON_BIG, ICON_SMALL
                    handle = user32.LoadImageW(None, ICON_PATH, 1, size, size, 0x10)
                    user32.SendMessageW(hwnd, 0x80, kind, handle)  # WM_SETICON
                return
            time.sleep(0.5)
    threading.Thread(target=run, daemon=True).start()


def open_window(url: str):
    """Tries pywebview (native window); falls back to Edge in app mode (no address bar)."""
    try:
        import webview  # pip install pywebview
        win = webview.create_window("Live Translator", url, width=1040, height=760,
                                    min_size=(800, 580))
        win.events.closed += hard_exit
        webview.start()
        return
    except Exception as e:  # noqa: BLE001
        print(f"pywebview unavailable ({type(e).__name__}), opening an Edge window instead...")
    try:
        subprocess.Popen(["cmd", "/c", "start", "", "msedge", f"--app={url}",
                          "--window-size=1040,760"])
    except Exception:  # noqa: BLE001
        webbrowser.open(url)
    wait_for_window_close()


def main():
    threading.Thread(target=segmenter, daemon=True).start()
    threading.Thread(target=worker, daemon=True).start()
    threading.Thread(target=broadcaster, daemon=True).start()
    threading.Thread(target=watchdog, daemon=True).start()
    srv, port = start_server()
    url = f"http://127.0.0.1:{port}/"
    stream = sd.InputStream(samplerate=SR, channels=1, dtype="float32",
                            blocksize=BLOCK, device=MIC_DEVICE, callback=audio_callback)
    stream.start()
    print(f"Interface: {url}  (Ctrl+C to quit)", flush=True)
    apply_window_icon()
    try:
        open_window(url)
    except KeyboardInterrupt:
        pass
    finally:
        shutdown.set()
        listening.clear()

        def cleanup():
            for fn in (stream.abort, stream.close):
                try:
                    fn()
                except Exception:  # noqa: BLE001
                    pass
        t = threading.Thread(target=cleanup, daemon=True)
        t.start()
        t.join(2)  # if the audio thread hangs, don't wait longer than 2 s
        hard_exit()


if __name__ == "__main__":
    main()